"""Read-structure survey for runs without a sample sheet or primer list.

For Nextera-style libraries ([P5][i5][A1] insert [A2][i7][P7]) it counts index
combinations, clusters inserts by their two ends to find the amplicons, infers
primer sequences from the end consensus, flags unconverted products and places
the bisulfite products on GRCm39 with the in-silico PCR used by the analysis.
"""
from __future__ import annotations

from collections import Counter, defaultdict
import json
import multiprocessing
import statistics

import edlib
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

from . import bsmap, indexes
from .progress import emit, log
from .reference import find_amplicon, revcomp

P5_TAIL = "GATCTACAC"
P7_TAIL = "ATCTCGTATGCCGTCTTCTGCTTG"
A1, A2_RC = indexes.A1, revcomp(indexes.A2)
END = 20                      # bases compared when clustering insert ends
PRIMER = 22                   # longest primer reported
NONE = "なし"
STRUCTURES = {"both": "両端にNexteraアダプター", "a1": "A1（i5側）のみ", "a2": "A2（i7側）のみ", "none": "アダプターなし"}


def _find(pattern, sequence, k):
    hit = edlib.align(pattern, sequence, mode="HW", task="locations", k=k)
    if hit["editDistance"] < 0:
        return None
    start, end = hit["locations"][0]
    if start is None:           # edlib may omit the start; estimate it from the pattern length
        start = max(0, end + 1 - len(pattern))
    return start, end + 1, hit["editDistance"]


def _name(sequence, table):
    """Kit index within one edit of the sequence read, else the raw sequence."""
    if not sequence:
        return NONE
    scored = sorted((edlib.align(value, sequence, mode="NW")["editDistance"], name) for name, value in table.items())
    if scored[0][0] <= 1 and (len(scored) == 1 or scored[1][0] > scored[0][0]):
        return scored[0][1]
    return sequence if len(sequence) == 8 else "?"


def survey(sequence):
    """(structure, i5, i7, insert) for one read, oriented so that A1 comes first."""
    sequence = sequence.upper()
    if len(sequence) < 60:
        return "none", NONE, NONE, None
    best = None
    for oriented in (sequence, revcomp(sequence)):
        a1, a2 = _find(A1, oriented, 7), _find(A2_RC, oriented, 7)
        score = sum(hit[2] if hit else 20 for hit in (a1, a2))
        if best is None or score < best[0]:
            best = (score, oriented, a1, a2)
    _, read, a1, a2 = best
    if a1 and a2 and a2[0] < a1[1]:
        a2 = None
    structure = "both" if a1 and a2 else "a1" if a1 else "a2" if a2 else "none"
    i5 = i7 = NONE
    if a1:
        window = read[max(0, a1[0] - 30):a1[0]]
        tail = _find(P5_TAIL, window, 2)
        if tail and 6 <= len(window) - tail[1] <= 10:
            i5 = _name(window[tail[1]:], indexes.I5)
    if a2:
        window = read[a2[1]:a2[1] + 40]
        tail = _find(P7_TAIL, window, 3)
        if tail and 6 <= tail[0] <= 10:
            i7 = _name(window[:tail[0]], indexes.I7)        # read after A2 = sample-sheet orientation
    insert = read[a1[1]:a2[0]] if structure == "both" else None
    return structure, i5, i7, insert


def _survey_chunk(chunk):
    return [survey(sequence) for sequence in chunk]


def _consensus(sequences, bisulfite_base):
    """Primer-like consensus: constant columns from the 5' end, C/T (or G/A) mixes as IUPAC."""
    result = []
    for column in range(PRIMER):
        counts = Counter(s[column] for s in sequences if len(s) > column)
        total = sum(counts.values())
        if not total:
            break
        base, top = counts.most_common(1)[0]
        mix = {"C": ("C", "T", "Y"), "G": ("G", "A", "R")}[bisulfite_base]
        if top >= 0.85 * total:
            result.append(base)
        elif counts[mix[0]] + counts[mix[1]] >= 0.85 * total and min(counts[mix[0]], counts[mix[1]]) >= 0.1 * total:
            result.append(mix[2])
        else:
            break
    return "".join(result)


def _nonconversion(insert):
    """Fraction of non-CpG C in an insert: ~0 after bisulfite, ~0.2 without."""
    cs = sum(1 for i, base in enumerate(insert) if base == "C" and insert[i + 1:i + 2] != "G")
    return cs / max(1, len(insert))


def cluster_inserts(inserts, minimum):
    """Groups of inserts sharing both ends; near-identical end pairs are merged."""
    exact = Counter((s[:END], s[-END:]) for s in inserts if len(s) >= 2 * END)
    clusters = []
    for (left, right), count in exact.most_common(3000):
        if count < 5:
            break
        for cluster in clusters:
            if (edlib.align(left, cluster["left"], mode="NW")["editDistance"] <= 2
                    and edlib.align(right, cluster["right"], mode="NW")["editDistance"] <= 2):
                cluster["keys"].add((left, right))
                break
        else:
            clusters.append(dict(left=left, right=right, keys={(left, right)}))
    members = defaultdict(list)
    lookup = {key: n for n, cluster in enumerate(clusters) for key in cluster["keys"]}
    for i, s in enumerate(inserts):
        n = lookup.get((s[:END], s[-END:]))
        if n is not None:
            members[n].append(i)
    return [(clusters[n], members[n]) for n in sorted(members, key=lambda n: -len(members[n]))
            if len(members[n]) >= minimum]


def _map_locus(mapper, inserts, sample=40):
    """(chrom, start, end, strand of X) from the longest inserts of a cluster.

    The longest inserts are the full-length products; shorter ones are often
    primer-to-primer by-products. Strand "+" means the top strand was amplified.
    """
    placed = Counter()
    spans = defaultdict(list)
    for insert in sorted(inserts, key=len, reverse=True)[:sample]:
        hits = mapper.map(insert)
        if hits:
            hit = hits[0]
            key = (hit.chrom, "+" if hit.top else "-", hit.start // 1000)
            placed[key] += 1
            spans[key].append((hit.start, hit.end))
    if not placed:
        return None
    key, count = placed.most_common(1)[0]
    if count < 0.5 * sum(placed.values()):
        return None
    starts, ends = zip(*spans[key])
    return key[0], int(statistics.median(starts)), int(statistics.median(ends)), key[1]


def run(out, request, files, resources, read_records):
    emit("explore", message=f"{len(files)}ファイルのリード構造を調べています。")
    limit = request.get("max_reads") or 0
    sequences = []
    for _, path in files:
        for _, sequence, _ in read_records(path):
            sequences.append(sequence)
            if limit and len(sequences) >= limit:
                break
        if limit and len(sequences) >= limit:
            break
    chunks = [sequences[i:i + 1000] for i in range(0, len(sequences), 1000)]
    results = []
    with multiprocessing.get_context("fork").Pool(request["threads"]) as pool:
        for part in pool.imap(_survey_chunk, chunks):
            results.extend(part)
            if len(results) % 50000 < 1000:
                emit("explore", done=len(results), message=f"{len(results):,}リードを調べました。")
    total = len(results)
    structure = Counter(r[0] for r in results)
    combos = Counter((r[1], r[2]) for r in results if r[0] != "none")
    inserts, insert_index = [], []
    for n, r in enumerate(results):
        if r[3]:
            inserts.append(r[3])
            insert_index.append(n)
    minimum = max(50, int(0.005 * len(inserts)))
    mapper = None
    if bsmap.index_ready(resources):
        emit("explore", message="ゲノム索引を読み込んでいます。")
        mapper = bsmap.BisulfiteMapper(bsmap.index_path(resources))
    candidates = []
    for number, (cluster, members) in enumerate(cluster_inserts(inserts, minimum), 1):
        seqs = [inserts[i] for i in members]
        lengths = [len(s) for s in seqs]
        unconverted = statistics.median(_nonconversion(s) for s in seqs)
        converted = unconverted < 0.03
        forward = _consensus(seqs, "C")
        reverse = _consensus([revcomp(s) for s in seqs], "G" if converted else "C")
        peaks = Counter(round(length, -1) for length in lengths).most_common(3)
        entry = dict(number=number, reads=len(seqs), fraction=round(100.0 * len(seqs) / max(1, len(inserts)), 1),
                     converted=converted, nonconversion=round(100 * unconverted, 1), forward=forward, reverse=reverse,
                     median_length=int(statistics.median(lengths)), peaks=peaks, name=f"候補{number}",
                     chrom=None, start=None, end=None, strand=None, size=None, full_length=None, note="")
        combo = Counter((results[insert_index[i]][1], results[insert_index[i]][2]) for i in members)
        entry["combos"] = [dict(i5=a, i7=b, reads=c) for (a, b), c in combo.most_common()
                           if c >= max(10, 0.005 * len(seqs))]
        locus = _map_locus(mapper, seqs) if mapper else None
        if locus:
            chrom, lo, hi, strand = locus
            entry.update(chrom=chrom, start=lo + 1, end=hi, strand=strand if converted else None, size=hi - lo)
        if not converted:
            entry["note"] = "未変換（バイサルファイト処理されていない別の産物）"
        elif len(forward) >= 15 and len(reverse) >= 15:
            # The primer search gives exact product ends where the inferred primers match the genome.
            try:
                product = find_amplicon(resources, f"候補{number}", forward, reverse, max(5000, 2 * (entry["size"] or 0)))
                if not locus or (product["chrom"] == locus[0] and min(product["end"], locus[2]) > max(product["start"], locus[1])):
                    entry.update(chrom=product["chrom"], start=product["start"] + 1, end=product["end"],
                                 strand=product["strand"], size=product["end"] - product["start"])
                    if product.get("swapped"):
                        entry["forward"], entry["reverse"] = product["forward"], product["reverse"]
            except ValueError as error:
                if not locus:
                    entry["note"] = str(error)
        elif not locus:
            entry["note"] = "末端がそろわずプライマーを推定できません"
        if entry["chrom"]:
            size = entry["size"]
            full = sum(1 for length in lengths if abs(length - size) <= 0.1 * size)
            entry["full_length"] = round(100.0 * full / len(seqs), 1)
            entry["name"] = f"chr{entry['chrom']}_{entry['start']}_{size}bp"
            if converted and full < 0.5 * len(seqs):
                entry["note"] = (f"完全長（{size} bp）のリードは{100.0 * full / len(seqs):.0f}%のみ。"
                                 "残りは両プライマーを持つ短い副産物の可能性があります")
        candidates.append(entry)
        log(f"候補{number}: {len(seqs):,}リード F={forward} R={reverse} {entry['note']}")
    result = dict(mode="explore", info=dict(files=[relative for relative, _ in files], reads=total,
                                            inserts=len(inserts), limited=bool(limit and total >= limit)),
                  structure=[dict(構造=STRUCTURES[key], リード数=structure[key],
                                  割合=round(100.0 * structure[key] / max(1, total), 1)) for key in STRUCTURES],
                  combos=[dict(i5=a, i7=b, reads=c) for (a, b), c in combos.most_common() if c >= 10],
                  candidates=candidates, warnings=[])
    if structure["both"] + structure["a1"] + structure["a2"] < 0.2 * total:
        result["warnings"].append("Nexteraアダプターを持つリードが少ないため、インラインタグのライブラリの可能性があります。"
                                  "この調査はNextera形式のライブラリ向けです。")
    (out / "results.json").write_text(json.dumps(result, ensure_ascii=False), encoding="utf-8")
    _workbook(out / "explore.xlsx", result)


def _workbook(path, result):
    workbook = Workbook()
    header = PatternFill(start_color="DDEBF7", end_color="DDEBF7", fill_type="solid")

    def sheet(title, columns, rows):
        ws = workbook.create_sheet(title)
        ws.append(columns)
        for cell in ws[1]:
            cell.font, cell.fill = Font(bold=True), header
        for row in rows:
            ws.append(row)
        for column in ws.columns:
            width = max(len(str(cell.value)) for cell in column if cell.value is not None)
            ws.column_dimensions[column[0].column_letter].width = min(60, max(6, width + 2))
    workbook.remove(workbook.active)
    sheet("構造", ["構造", "リード数", "割合(%)"], [list(row.values()) for row in result["structure"]])
    sheet("アンプリコン候補", ["名前", "リード数", "インサート中の割合(%)", "変換", "非CpGのC(%)", "Forward推定", "Reverse推定",
                        "位置", "鎖", "産物長", "完全長リード(%)", "インサート長の中央値", "主なインサート長", "メモ"],
          [[c["name"], c["reads"], c["fraction"], "変換済み" if c["converted"] else "未変換", c["nonconversion"],
            c["forward"], c["reverse"], f"chr{c['chrom']}:{c['start']}-{c['end']}" if c["chrom"] else "",
            c["strand"] or "", c["size"] or "", c["full_length"] if c["full_length"] is not None else "",
            c["median_length"], ", ".join(f"{length}({count})" for length, count in c["peaks"]), c["note"]]
           for c in result["candidates"]])
    i5s = sorted({row["i5"] for row in result["combos"]})
    i7s = sorted({row["i7"] for row in result["combos"]})
    counts = {(row["i5"], row["i7"]): row["reads"] for row in result["combos"]}
    sheet("インデックス組み合わせ", ["i5 \\ i7"] + i7s, [[a] + [counts.get((a, b), "") for b in i7s] for a in i5s])
    sheet("候補別インデックス", ["候補", "i5", "i7", "リード数"],
          [[c["name"], row["i5"], row["i7"], row["reads"]] for c in result["candidates"] for row in c["combos"]])
    workbook.save(path)
