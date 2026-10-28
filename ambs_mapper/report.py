"""Tallies, QUMA-san style photos, charts and Excel/CSV outputs."""
from __future__ import annotations

from collections import Counter, defaultdict
import csv
import json
from pathlib import Path
import random
import re

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt  # noqa: E402
from matplotlib.colors import ListedColormap  # noqa: E402
from openpyxl import Workbook  # noqa: E402
from openpyxl.styles import Alignment, Font, PatternFill  # noqa: E402

from .caller import ALLELE_LABELS, STATUS_LABELS  # noqa: E402
from .indexes import LAYOUTS  # noqa: E402

plt.rcParams["font.family"] = ["IPAGothic", "DejaVu Sans"]
STATUS_ORDER = ["pass", "low_coverage", "low_identity", "high_mismatch", "low_conversion", "unconverted_cph"]
READ_STATUSES = ["unmapped", "chimera", "tag_unknown", "tag_ambiguous"]
ALLELES = ["mother", "father", "unassigned", "same"]
BLACK = PatternFill(start_color="000000", end_color="000000", fill_type="solid")
GREY = PatternFill(start_color="D9D9D9", end_color="D9D9D9", fill_type="solid")
HEADER = PatternFill(start_color="DDEBF7", end_color="DDEBF7", fill_type="solid")


class Collector:
    """Accumulates per-read records without keeping failed reads in memory."""

    def __init__(self, samples, targets):
        self.samples, self.targets = samples, {t.name: t for t in targets}
        self.sample_status = defaultdict(Counter)       # sample index -> status counts
        self.amplicon_status = defaultdict(Counter)     # amplicon -> reads not assigned to a sample
        self.patterns = defaultdict(list)               # (sample index, allele) -> patterns
        self.source_status = defaultdict(Counter)       # FASTQ -> status counts
        self.single = Counter()                         # sample index -> reads with one index only
        self.off_target = Counter()                     # (chrom, 1-kb bin, bisulfite strand) -> reads
        self.passed = 0

    def add(self, record):
        status = record["status"]
        self.source_status[record["source"]][status] += 1
        if status in ("off_target", "wrong_strand"):
            self.off_target[(record["chrom"], record["gstart"] // 1000 * 1000, record["bs_strand"])] += 1
        if record["sample"] < 0:
            self.amplicon_status[record["amplicon"] or "-"][status] += 1
            return
        self.sample_status[record["sample"]][status] += 1
        self.single[record["sample"]] += record.get("tag_sides") == 1
        if status == "pass":
            self.passed += 1
            self.patterns[(record["sample"], record["allele"])].append(record["pattern"])

    def drop(self, sample):
        """Move an automatic sample's reads to the amplicon's unassigned counts."""
        counts = self.sample_status.pop(sample.index, Counter())
        self.amplicon_status[sample.amplicon]["tag_minor"] += sum(counts.values())
        for allele in ALLELES:
            self.patterns.pop((sample.index, allele), None)


def level(patterns, site=None):
    """(methylated, called) over all CpGs, or at one site."""
    methylated = called = 0
    for pattern in patterns:
        values = pattern if site is None else pattern[site]
        methylated += values.count("M")
        called += values.count("M") + values.count("U")
    return methylated, called


def percent(methylated, called):
    return round(100.0 * methylated / called, 2) if called else None


def quma_order(patterns):
    """QUMA-san ordering: fewer methylated CpGs first, then by binary pattern."""
    def key(pattern):
        return (pattern.count("M"), sum(1 << i for i, value in enumerate(pattern) if value == "M"))
    return sorted(patterns, key=key)


def safe(text):
    return re.sub(r"[^0-9A-Za-z._-]+", "_", text).strip("_")[:50] or "x"


def groups(samples, collector):
    """Display groups: each sample x allele that has passed reads, plus the all-allele total."""
    result = []
    for sample in samples:
        alleles = ["same"] if sample.mother == sample.father else ["mother", "father", "unassigned"]
        for allele in alleles + (["all"] if sample.mother != sample.father else []):
            if allele == "all":
                patterns = [p for a in ("mother", "father", "unassigned") for p in collector.patterns[(sample.index, a)]]
            else:
                patterns = collector.patterns[(sample.index, allele)]
            label = {"all": "全リード"}.get(allele) or ALLELE_LABELS[allele]
            strain = {"mother": sample.mother, "father": sample.father, "same": sample.mother}.get(allele, "")
            result.append(dict(sample=sample, allele=allele, label=label, strain=strain, patterns=patterns))
    return result


def draw_photo(path, patterns, positions, title, max_rows, seed=0):
    total = len(patterns)
    shown = patterns if total <= max_rows else random.Random(seed).sample(patterns, max_rows)
    shown = quma_order(shown)
    columns = len(positions)
    rows = len(shown)
    width = min(18, max(4, 0.32 * columns + 1.5))
    if rows <= 60:
        height = max(2.2, 0.22 * rows + 1.4)
        figure, axis = plt.subplots(figsize=(width, height))
        for row, pattern in enumerate(shown):
            y = rows - row
            axis.plot([0.5, columns + 0.5], [y, y], color="#888888", linewidth=0.6, zorder=1)
            for column, value in enumerate(pattern, 1):
                if value == "M":
                    axis.scatter(column, y, s=60, color="black", edgecolors="black", zorder=2)
                elif value == "U":
                    axis.scatter(column, y, s=60, color="white", edgecolors="black", zorder=2)
                elif value in "-N":
                    axis.scatter(column, y, s=12, color="#bbbbbb", zorder=2)
        axis.set_ylim(0.3, rows + 0.7)
        axis.set_yticks([])
    else:
        height = min(16, max(3, 0.03 * rows + 1.6))
        figure, axis = plt.subplots(figsize=(width, height))
        values = {"U": 0, "M": 1, "-": 2, "N": 2, ".": 3}
        matrix = [[values[value] for value in pattern] for pattern in shown]
        axis.imshow(matrix, aspect="auto", interpolation="nearest",
                    cmap=ListedColormap(["#ffffff", "#000000", "#d0d0d0", "#f2e3c9"]), vmin=0, vmax=3,
                    extent=(0.5, columns + 0.5, rows + 0.5, 0.5))
        axis.set_ylabel("リード")
    axis.set_xlim(0.5, columns + 0.5)
    step = max(1, columns // 30)
    ticks = list(range(1, columns + 1, step))
    axis.set_xticks(ticks)
    axis.set_xticklabels([str(t) for t in ticks], fontsize=7)
    axis.set_xlabel("CpG番号")
    for spine in ("top", "right", "left"):
        axis.spines[spine].set_visible(False)
    methylated, called = level(patterns)
    note = f"表示 {rows:,}/{total:,}リード" + ("（無作為抽出）" if total > rows else "")
    axis.set_title(f"{title}\nメチル化率 {percent(methylated, called) or 0:.1f}%  {note}", fontsize=9)
    figure.tight_layout()
    figure.savefig(path, dpi=150)
    plt.close(figure)


def draw_profiles(path, target_name, positions, entries):
    figure, axis = plt.subplots(figsize=(min(16, max(6, 0.3 * len(positions) + 3)), 4))
    for entry in entries:
        values = [percent(*level(entry["patterns"], i)) for i in range(len(positions))]
        axis.plot(range(1, len(positions) + 1), [v if v is not None else float("nan") for v in values],
                  marker="o", markersize=3, linewidth=1.2, label=f"{entry['sample'].name} {entry['label']}")
    axis.set_ylim(-3, 103)
    axis.set_xlabel("CpG番号")
    axis.set_ylabel("メチル化率 (%)")
    axis.set_title(f"{target_name}: CpG別メチル化率")
    axis.grid(alpha=0.3)
    if entries:
        axis.legend(fontsize=7, loc="center left", bbox_to_anchor=(1.01, 0.5))
    figure.tight_layout()
    figure.savefig(path, dpi=150)
    plt.close(figure)


def draw_means(path, target_name, entries):
    names = list(dict.fromkeys(entry["sample"].name for entry in entries))
    alleles = list(dict.fromkeys(entry["label"] for entry in entries))
    figure, axis = plt.subplots(figsize=(max(5, 0.9 * len(names) * max(1, len(alleles)) / 2 + 2), 4))
    width = 0.8 / max(1, len(alleles))
    for a, label in enumerate(alleles):
        xs, ys = [], []
        for n, name in enumerate(names):
            for entry in entries:
                if entry["sample"].name == name and entry["label"] == label:
                    value = percent(*level(entry["patterns"]))
                    xs.append(n + (a - (len(alleles) - 1) / 2) * width)
                    ys.append(value or 0)
        axis.bar(xs, ys, width=width, label=label)
    axis.set_xticks(range(len(names)))
    axis.set_xticklabels(names, rotation=30, ha="right", fontsize=8)
    axis.set_ylim(0, 100)
    axis.set_ylabel("メチル化率 (%)")
    axis.set_title(f"{target_name}: サンプル別メチル化率")
    axis.legend(fontsize=8)
    figure.tight_layout()
    figure.savefig(path, dpi=150)
    plt.close(figure)


def _header(sheet, values):
    sheet.append(values)
    for cell in sheet[sheet.max_row]:
        cell.font = Font(bold=True)
        cell.fill = HEADER


def _autosize(sheet, limit=40):
    for column in sheet.columns:
        width = max((len(str(cell.value)) for cell in column if cell.value is not None), default=4)
        sheet.column_dimensions[column[0].column_letter].width = min(limit, max(6, width + 2))


def _sheet_name(used, text):
    name = re.sub(r"[\[\]:*?/\\]", "_", text)[:31]
    base, number = name, 2
    while name in used:
        suffix = f"~{number}"
        name, number = base[:31 - len(suffix)] + suffix, number + 1
    used.add(name)
    return name


def sample_rows(samples, collector):
    rows = []
    for sample in samples:
        counts = collector.sample_status[sample.index]
        row = dict(サンプル=sample.name, 群=sample.group, アンプリコン=sample.amplicon, タグF=sample.labels[0],
                   タグR=sample.labels[1], 母系統=sample.mother, 父系統=sample.father, 割当リード=sum(counts.values()))
        row["うち片側タグのみ"] = collector.single[sample.index]
        for status in STATUS_ORDER:
            row[STATUS_LABELS[status]] = counts[status]
        for allele in ALLELES:
            applies = (allele == "same") == (sample.mother == sample.father)
            patterns = collector.patterns[(sample.index, allele)]
            row[f"{ALLELE_LABELS[allele]}リード"] = len(patterns) if applies else None
            row[f"{ALLELE_LABELS[allele]}メチル化率(%)"] = percent(*level(patterns)) if applies else None
        everything = [p for a in ALLELES for p in collector.patterns[(sample.index, a)]]
        row["全合格リードのメチル化率(%)"] = percent(*level(everything))
        rows.append(row)
    return rows


def cpg_rows(entries, details_by_name):
    rows = []
    for entry in entries:
        detail = details_by_name[entry["sample"].amplicon]
        for i, position in enumerate(detail["site_positions"]):
            methylated, called = level(entry["patterns"], i)
            rows.append(dict(アンプリコン=detail["name"], サンプル=entry["sample"].name, 群=entry["sample"].group,
                             アレル=entry["label"], 系統=entry["strain"], CpG番号=i + 1,
                             位置=f"{detail['chrom']}:{position}", メチル化=methylated, 判定数=called,
                             メチル化率=percent(methylated, called)))
    return rows


def _write_csv(path, rows):
    if not rows:
        path.write_text("", encoding="utf-8-sig")
        return
    with open(path, "w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.DictWriter(handle, list(rows[0]))
        writer.writeheader()
        writer.writerows(rows)


def _settings_rows(request, info, warnings):
    p = request["params"]
    rows = [("解析名", request.get("title") or ""), ("参照", "GRCm39 / MGP REL-2112 v8"),
            ("入力ファイル", "\n".join(info.get("files", []))), ("総リード数", info.get("reads")),
            ("リード処理時間(秒)", info.get("seconds")),
            ("アンプリコン被覆率下限(%)", p["min_coverage"]), ("一致率下限(%) [QUMA -p]", p["min_identity"]),
            ("ミスマッチ上限(%) [QUMA -m]", p["max_mismatch"]), ("CpH変換率下限(%) [QUMA -c]", p["min_conversion"]),
            ("未変換CpH上限(個) [QUMA -u]", "制限なし" if p["max_unconverted"] is None else p["max_unconverted"]),
            ("タグ形式", LAYOUTS[request.get("layout", "inline")]),
            ("自動サンプル", "なし" if not request.get("auto") else
             f"母 {request['auto']['mother']} / 父 {request['auto']['father']}（合格リード{p['auto_min_reads']}以上の組み合わせ）"),
            ("タグ許容誤差", p["tag_max_errors"]), ("タグとプライマー間の許容塩基", p["tag_slack"]),
            ("SNP塩基品質下限", p["snp_min_quality"]), ("アレル判定の一致率下限(%)", p["allele_min_fraction"]),
            ("プライマー内CpGを除外", "はい" if p["exclude_primer_cpg"] else "いいえ")]
    rows += [("注意", text) for text in warnings]
    return rows


def write_photo_workbook(path, entries, details_by_name, limit=10000):
    workbook = Workbook()
    workbook.remove(workbook.active)
    used = set()
    for entry in entries:
        if not entry["patterns"]:
            continue
        detail = details_by_name[entry["sample"].amplicon]
        sheet = workbook.create_sheet(_sheet_name(used, f"{entry['sample'].name}_{entry['label']}"))
        sheet.append([f"{detail['name']} / {entry['sample'].name} / {entry['label']} {entry['strain']}",
                      f"合格リード {len(entry['patterns']):,}" + (f"（先頭{limit:,}件を表示）" if len(entry['patterns']) > limit else "")])
        _header(sheet, ["リード"] + [f"CpG{i + 1}" for i in range(len(detail["sites"]))] + ["メチル化数", "メチル化率(%)"])
        for number, pattern in enumerate(quma_order(entry["patterns"])[:limit], 1):
            methylated, called = level([pattern])
            sheet.append([number] + [{"M": "●", "U": "", "-": "-", "N": "x", ".": "/"}[v] for v in pattern]
                         + [methylated, percent(methylated, called)])
            for cell, value in zip(sheet[sheet.max_row][1:], pattern):
                if value == "M":
                    cell.fill = BLACK
                elif value in "-N.":
                    cell.fill = GREY
        for column in range(2, len(detail["sites"]) + 2):
            sheet.column_dimensions[sheet.cell(row=2, column=column).column_letter].width = 4.5
        sheet.freeze_panes = "B3"
    if not workbook.sheetnames:
        workbook.create_sheet("合格リードなし")
    workbook.save(path)


BS_STRANDS = {"OT": "上鎖", "CTOT": "上鎖（相補）", "OB": "下鎖", "CTOB": "下鎖（相補）"}


def off_target_rows(collector, limit=50):
    """Most frequent loci outside the amplicons (1-kb bins), e.g. mispriming or other products."""
    return [dict(領域=f"chr{chrom}:{start + 1:,}-{start + 1000:,}", 由来鎖=BS_STRANDS.get(strand, strand), リード数=count)
            for (chrom, start, strand), count in collector.off_target.most_common(limit)]


def write_check(out, request, samples, details, warnings):
    (out / "results.json").write_text(json.dumps(dict(mode="check", warnings=warnings, amplicons=[
        {k: v for k, v in d.items() if k != "sites"} for d in details], samples=[], files=[]),
        ensure_ascii=False), encoding="utf-8")


def write_all(out, request, samples, targets, details, warnings, collector, info):
    details_by_name = {d["name"]: d for d in details}
    entries = groups(samples, collector)
    photos = out / "photos"
    photos.mkdir(exist_ok=True)
    max_rows = max(1, request["params"]["photo_rows"])
    images = []
    for entry in entries:
        entry["image"] = None
        if not entry["patterns"]:
            continue
        detail = details_by_name[entry["sample"].amplicon]
        name = f"photos/{safe(detail['name'])}__{safe(entry['sample'].name)}__{entry['allele']}.png"
        title = f"{detail['name']} / {entry['sample'].name} / {entry['label']} {entry['strain']}".strip()
        draw_photo(out / name, entry["patterns"], detail["sites"], title, max_rows)
        entry["image"] = name
    for detail in details:
        members = [e for e in entries if e["sample"].amplicon == detail["name"] and e["allele"] != "all"
                   and e["patterns"]]
        profile = f"plots/{safe(detail['name'])}__cpg.png"
        means = f"plots/{safe(detail['name'])}__mean.png"
        (out / "plots").mkdir(exist_ok=True)
        if detail["sites"]:
            draw_profiles(out / profile, detail["name"], detail["sites"], members)
            draw_means(out / means, detail["name"], members)
            images.append(dict(amplicon=detail["name"], profile=profile, means=means))

    summary = sample_rows(samples, collector)
    cpgs = cpg_rows(entries, details_by_name)
    _write_csv(out / "samples.csv", summary)
    _write_csv(out / "cpg_methylation.csv", cpgs)
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "設定"
    for row in _settings_rows(request, info, warnings):
        sheet.append(list(row))
    for cell in sheet["B"]:
        cell.alignment = Alignment(wrap_text=True, vertical="top")
    sheet.column_dimensions["A"].width, sheet.column_dimensions["B"].width = 34, 90
    for title, rows in (("サンプル別", summary), ("CpG別", cpgs)):
        sheet = workbook.create_sheet(title)
        if rows:
            _header(sheet, list(rows[0]))
            for row in rows:
                sheet.append(list(row.values()))
            _autosize(sheet)
    sheet = workbook.create_sheet("リード振り分け")
    _header(sheet, ["区分", "対象", "リード数"])
    for amplicon, counts in sorted(collector.amplicon_status.items()):
        for status, count in counts.most_common():
            sheet.append([f"未割当: {STATUS_LABELS[status]}", amplicon, count])
    for source, counts in sorted(collector.source_status.items()):
        sheet.append(["FASTQ合計", source, sum(counts.values())])
    _autosize(sheet, 60)
    sheet = workbook.create_sheet("アンプリコン外")
    _header(sheet, ["領域(1 kb単位)", "由来鎖", "リード数"])
    for row in off_target_rows(collector, 500):
        sheet.append(list(row.values()))
    _autosize(sheet)
    sheet = workbook.create_sheet("アンプリコン")
    _header(sheet, ["アンプリコン", "位置", "鎖", "長さ", "Forward", "Reverse", "CpG数", "母系統", "父系統",
                    "判定に使えるSNP", "C/T型で使えないSNP", "indel", "配列(変換前、Forward側の鎖)"])
    for detail in details:
        for pair in detail["pairs"] or [dict(mother="", father="", informative=0, snps=[], indels=[])]:
            sheet.append([detail["name"], f"{detail['chrom']}:{detail['start']}-{detail['end']}", detail["strand"],
                          detail["size"], detail["forward"], detail["reverse"], len(detail["sites"]),
                          pair["mother"], pair["father"], pair["informative"],
                          sum(1 for s in pair["snps"] if not s["informative"]), len(pair["indels"]), detail["sequence"]])
    _autosize(sheet, 30)
    sheet = workbook.create_sheet("SNP")
    _header(sheet, ["アンプリコン", "ゲノム位置", "アンプリコン内位置", "母系統", "母塩基(Forward鎖)", "父系統",
                    "父塩基(Forward鎖)", "アレル判定に使用"])
    for detail in details:
        for pair in detail["pairs"]:
            for snp in pair["snps"]:
                sheet.append([detail["name"], f"{detail['chrom']}:{snp['genome']}", snp["x"] + 1, pair["mother"],
                              snp["mother"], pair["father"], snp["father"],
                              "はい" if snp["informative"] else "いいえ（C/T型）"])
    _autosize(sheet)
    workbook.save(out / "summary.xlsx")
    write_photo_workbook(out / "photo.xlsx", entries, details_by_name)

    result = dict(
        mode="analyze", info=info, warnings=warnings, images=images, samples=summary,
        amplicons=[{k: v for k, v in d.items() if k not in ("sites", "sequence")} for d in details],
        unassigned={amplicon: {STATUS_LABELS[s]: c for s, c in counts.items()}
                    for amplicon, counts in collector.amplicon_status.items()},
        off_target=off_target_rows(collector),
        groups=[dict(sample=e["sample"].name, group=e["sample"].group, amplicon=e["sample"].amplicon,
                     allele=e["label"], strain=e["strain"], reads=len(e["patterns"]),
                     methylation=percent(*level(e["patterns"])), image=e["image"]) for e in entries])
    (out / "results.json").write_text(json.dumps(result, ensure_ascii=False), encoding="utf-8")
